"""Depodaki veriyle uçtan uca: README'deki bulgular yeniden üretilebiliyor mu?"""
import pandas as pd
import pytest
from openpyxl import load_workbook

import analysis
from analysis import OEE_TARGET, add_oee, decorator_crew_test, oee_by, spc_table


@pytest.fixture(scope="module")
def data():
    shifts, events, quality = analysis.load()
    return add_oee(shifts), events, quality


def test_hat2_oee_hedefin_altinda(data):
    shifts, _, _ = data
    oee = oee_by(shifts, ["line"]).set_index("line")["oee"]
    assert oee["Line 2"] < OEE_TARGET
    assert oee["Line 1"] > oee["Line 2"]


def test_dekorator_durusu_crew_c_ye_ozgu(data):
    shifts, events, _ = data
    t = decorator_crew_test(shifts, events)
    assert t["p_value"] < 0.001
    assert t["by_crew"].idxmax() == "Crew C"
    assert t["crew_c_mean"] > 3 * t["others_mean"]


def test_spc_spek_disindan_once_uyari_verir(data):
    _, _, quality = data
    r = spc_table(quality, "Line 2")
    assert r["first_alert"] is not None
    assert r["first_alert"] < r["first_out_of_spec"]
    assert r["cpk_base"] > 1.33 > 1.0 > r["cpk_recent"]


def test_hat1_dome_basinci_kararli(data):
    _, _, quality = data
    r = spc_table(quality, "Line 1")
    assert pd.isna(r["first_out_of_spec"])
    assert r["cpk_recent"] > 1.33


def test_gunluk_rapor_sayfalari(data, tmp_path):
    shifts, events, quality = data
    spc = [spc_table(quality, l) for l in ["Line 1", "Line 2"]]
    path = tmp_path / "report.xlsx"
    day, kpi, losses = analysis.daily_report(shifts, events, spc, path)
    wb = load_workbook(path)
    assert wb.sheetnames == ["Daily KPI", "Top Losses", "SPC Status", "Action Log (PDCA)"]
    assert wb["Daily KPI"]["A1"].value.startswith("Daily Operations Meeting")
    assert day == shifts["date"].max()
    assert set(kpi.loc[kpi["shift"] == "DAY TOTAL", "line"]) == {"Line 1", "Line 2"}
