from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from scipy import stats

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
OUT = ROOT / "outputs"
CHARTS = OUT / "charts"

OEE_TARGET = 0.80
DOME_LSL = 90.0  # psi
SPC_BASELINE_DAYS = 60

C_LINE1, C_LINE2 = "#1f4e79", "#e07b39"
C_TARGET, C_ALERT, C_GREY = "#2e7d32", "#c62828", "#8c8c8c"

plt.rcParams.update({
    "figure.dpi": 130, "axes.spines.top": False, "axes.spines.right": False,
    "axes.grid": True, "grid.alpha": 0.25, "font.size": 10,
})


def load():
    shifts = pd.read_csv(DATA / "production_shifts.csv", parse_dates=["date"])
    events = pd.read_csv(DATA / "downtime_events.csv", parse_dates=["date"])
    quality = pd.read_csv(DATA / "quality_samples.csv", parse_dates=["date"])
    return shifts, events, quality


def add_oee(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df["availability"] = df["run_min"] / df["planned_min"]
    df["performance"] = df["total_cans"] / (df["run_min"] * df["ideal_cpm"])
    df["quality"] = df["good_cans"] / df["total_cans"]
    df["oee"] = df["availability"] * df["performance"] * df["quality"]
    return df


def oee_by(shifts: pd.DataFrame, keys) -> pd.DataFrame:
    g = shifts.assign(ideal_cans=shifts["run_min"] * shifts["ideal_cpm"]).groupby(keys).agg(
        planned_min=("planned_min", "sum"), run_min=("run_min", "sum"),
        ideal_cans=("ideal_cans", "sum"), total_cans=("total_cans", "sum"),
        good_cans=("good_cans", "sum"),
    )
    g["availability"] = g["run_min"] / g["planned_min"]
    g["performance"] = g["total_cans"] / g["ideal_cans"]
    g["quality"] = g["good_cans"] / g["total_cans"]
    g["oee"] = g["availability"] * g["performance"] * g["quality"]
    return g.reset_index()


def chart_weekly_oee(shifts):
    s = shifts.assign(week=shifts["date"].dt.to_period("W").dt.start_time)
    full_weeks = s.groupby("week")["date"].nunique().loc[lambda d: d == 7].index
    wk = oee_by(s[s["week"].isin(full_weeks)], ["week", "line"])
    fig, ax = plt.subplots(figsize=(9, 4))
    for line, color in [("Line 1", C_LINE1), ("Line 2", C_LINE2)]:
        d = wk[wk["line"] == line]
        ax.plot(d["week"], d["oee"] * 100, marker="o", color=color, label=line, lw=2)
    ax.axhline(OEE_TARGET * 100, color=C_TARGET, ls="--", lw=1.2, label=f"Target {OEE_TARGET:.0%}")
    ax.set_title("Weekly OEE by line", loc="left", fontweight="bold")
    ax.set_ylabel("OEE (%)")
    ax.legend(frameon=False, ncol=3)
    fig.tight_layout()
    fig.savefig(CHARTS / "01_weekly_oee.png")
    plt.close(fig)
    return wk


def chart_pareto(events):
    fig, axes = plt.subplots(1, 2, figsize=(12, 4.5), sharey=False)
    paretos = {}
    for ax, (line, color) in zip(axes, [("Line 1", C_LINE1), ("Line 2", C_LINE2)]):
        p = (events[events["line"] == line].groupby("reason")["minutes"].sum()
             .sort_values(ascending=False))
        cum = p.cumsum() / p.sum() * 100
        paretos[line] = pd.DataFrame({"minutes": p, "cum_pct": cum})
        ax.bar(range(len(p)), p.values / 60, color=color)
        ax.set_xticks(range(len(p)))
        ax.set_xticklabels(p.index, rotation=40, ha="right")
        ax.set_ylabel("Downtime (hours)")
        ax2 = ax.twinx()
        ax2.plot(range(len(p)), cum.values, color=C_GREY, marker="o", lw=1.5)
        ax2.axhline(80, color=C_GREY, ls=":", lw=1)
        ax2.set_ylim(0, 105)
        ax2.set_ylabel("Cumulative %")
        ax2.grid(False)
        ax.set_title(f"{line} - downtime Pareto", loc="left", fontweight="bold")
    fig.tight_layout()
    fig.savefig(CHARTS / "02_downtime_pareto.png")
    plt.close(fig)
    return paretos


def chart_crew_heatmap(shifts):
    crew = oee_by(shifts, ["line", "crew"]).pivot(index="line", columns="crew", values="oee")
    fig, ax = plt.subplots(figsize=(7, 2.8))
    im = ax.imshow(crew.values * 100, cmap="RdYlGn", vmin=70, vmax=85, aspect="auto")
    ax.set_xticks(range(len(crew.columns)), crew.columns)
    ax.set_yticks(range(len(crew.index)), crew.index)
    ax.grid(False)
    for i in range(crew.shape[0]):
        for j in range(crew.shape[1]):
            ax.text(j, i, f"{crew.iat[i, j]:.1%}", ha="center", va="center", fontweight="bold")
    fig.colorbar(im, ax=ax, label="OEE (%)")
    ax.set_title("OEE by line and crew", loc="left", fontweight="bold")
    fig.tight_layout()
    fig.savefig(CHARTS / "03_crew_heatmap.png")
    plt.close(fig)
    return crew


def chart_changeover(events):
    co = events[(events["line"] == "Line 1") & (events["reason"] == "Changeover")].copy()
    co["week"] = co["date"].dt.to_period("W").dt.start_time
    wk = co.groupby("week")["minutes"].agg(["sum", "mean", "count"]).reset_index()
    fig, ax = plt.subplots(figsize=(9, 3.8))
    ax.bar(wk["week"], wk["sum"] / 60, width=5, color=C_LINE1)
    ax.set_ylabel("Changeover hours / week")
    ax.set_title("Line 1 - weekly changeover time", loc="left", fontweight="bold")
    fig.tight_layout()
    fig.savefig(CHARTS / "04_line1_changeover.png")
    plt.close(fig)
    return co


A2, D3, D4, d2 = 0.577, 0.0, 2.114, 2.326  # n = 5


def spc_analysis(quality, line):
    q = quality[quality["line"] == line]
    sg = (q.groupby(["subgroup", "date"])["dome_reversal_psi"]
          .agg(xbar="mean", r=lambda s: s.max() - s.min()).reset_index()
          .sort_values("subgroup").reset_index(drop=True))

    start = sg["date"].min()
    base = sg[sg["date"] < start + pd.Timedelta(days=SPC_BASELINE_DAYS)]
    xbarbar, rbar = base["xbar"].mean(), base["r"].mean()
    ucl, lcl = xbarbar + A2 * rbar, xbarbar - A2 * rbar
    sigma = rbar / d2

    sg["rule1"] = (sg["xbar"] > ucl) | (sg["xbar"] < lcl)
    # 8 in a row on one side
    side = np.sign(sg["xbar"] - xbarbar)
    run = side.groupby((side != side.shift()).cumsum()).cumcount() + 1
    sg["rule_run8"] = run >= 8
    sg["alert"] = sg["rule1"] | sg["rule_run8"]

    recent = q[q["date"] >= q["date"].max() - pd.Timedelta(days=20)]["dome_reversal_psi"]
    cpk_base = (base["xbar"].mean() - DOME_LSL) / (3 * sigma)
    cpk_recent = (recent.mean() - DOME_LSL) / (3 * recent.std(ddof=1))

    fig, (a1, a2) = plt.subplots(2, 1, figsize=(11, 5.5), sharex=True,
                                 gridspec_kw={"height_ratios": [2, 1]})
    x = np.arange(len(sg))
    a1.plot(x, sg["xbar"], color=C_LINE2 if line == "Line 2" else C_LINE1, lw=1, marker=".", ms=3)
    a1.scatter(x[sg["alert"]], sg.loc[sg["alert"], "xbar"], color=C_ALERT, s=18, zorder=3,
               label="Out-of-control signal")
    for y, lab in [(ucl, "UCL"), (xbarbar, "CL"), (lcl, "LCL")]:
        a1.axhline(y, color=C_GREY, ls="--" if lab != "CL" else "-", lw=1)
        a1.text(len(sg) + 2, y, f"{lab} {y:.2f}", va="center", fontsize=8)
    a1.axhline(DOME_LSL, color=C_ALERT, lw=1.2, ls=":")
    a1.text(len(sg) + 2, DOME_LSL, f"LSL {DOME_LSL:.0f}", va="center", fontsize=8, color=C_ALERT)
    a1.axvline(len(base) - 0.5, color=C_GREY, lw=0.8)
    a1.set_ylabel("X-bar (psi)")
    a1.set_title(f"{line} - dome reversal pressure X-bar / R chart  "
                 f"(Cpk baseline {cpk_base:.2f} -> last 3 weeks {cpk_recent:.2f})",
                 loc="left", fontweight="bold")
    a1.legend(frameon=False, loc="lower left")
    a2.plot(x, sg["r"], color=C_GREY, lw=1)
    a2.axhline(rbar, color=C_GREY, lw=1)
    a2.axhline(D4 * rbar, color=C_GREY, ls="--", lw=1)
    a2.set_ylabel("Range (psi)")
    a2.set_xlabel("Subgroup (shift)")
    fig.tight_layout()
    fig.savefig(CHARTS / f"05_spc_{line.replace(' ', '').lower()}.png")
    plt.close(fig)

    after_base = sg["date"] >= start + pd.Timedelta(days=SPC_BASELINE_DAYS)
    first_alert = sg.loc[sg["alert"] & after_base & (sg["xbar"] < xbarbar), "date"]
    return {
        "line": line, "ucl": ucl, "cl": xbarbar, "lcl": lcl,
        "cpk_base": cpk_base, "cpk_recent": cpk_recent,
        "alerts": int(sg["alert"].sum()),
        "first_alert": first_alert.min() if len(first_alert) else None,
        "first_out_of_spec": q.loc[q["dome_reversal_psi"] < DOME_LSL, "date"].min(),
        "table": sg,
    }


def decorator_crew_test(shifts, events):
    l2 = shifts[shifts["line"] == "Line 2"][["date", "shift", "crew"]]
    dec = (events[(events["line"] == "Line 2") & (events["reason"] == "Decorator / printer")]
           .groupby(["date", "shift"])["minutes"].sum().reset_index())
    m = l2.merge(dec, on=["date", "shift"], how="left").fillna({"minutes": 0})
    c = m[m["crew"] == "Crew C"]["minutes"]
    others = m[m["crew"] != "Crew C"]["minutes"]
    u, p = stats.mannwhitneyu(c, others, alternative="greater")
    by_crew = m.groupby("crew")["minutes"].mean()
    return {"crew_c_mean": c.mean(), "others_mean": others.mean(), "p_value": p, "by_crew": by_crew}


HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
RED_FILL = PatternFill("solid", fgColor="F8D7DA")
GREEN_FILL = PatternFill("solid", fgColor="D4EDDA")


def style_sheet(ws, pct_cols=(), widths=None):
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col_idx in pct_cols:
        for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
            for c in row:
                c.number_format = "0.0%"
    for i, col in enumerate(ws.columns, start=1):
        width = (widths or {}).get(i) or max(len(str(c.value or "")) for c in col) + 3
        ws.column_dimensions[get_column_letter(i)].width = min(width, 45)
    ws.freeze_panes = "A2"


def daily_report(shifts, events, spc_results, path):
    day = shifts["date"].max()
    s_day = shifts[shifts["date"] == day]
    e_day = events[events["date"] == day]

    kpi = oee_by(s_day, ["line", "shift", "crew"])
    kpi_line = oee_by(s_day, ["line"]).assign(shift="DAY TOTAL", crew="")
    kpi = pd.concat([kpi, kpi_line], ignore_index=True)
    kpi["oee_target"] = OEE_TARGET
    kpi["gap_pts"] = (kpi["oee"] - OEE_TARGET) * 100
    kpi = kpi[["line", "shift", "crew", "availability", "performance", "quality",
               "oee", "oee_target", "gap_pts", "good_cans"]]

    losses = (e_day.groupby(["line", "reason"])
              .agg(stops=("event_id", "count"), minutes=("minutes", "sum"))
              .reset_index().sort_values("minutes", ascending=False))
    losses["lost_cans_est"] = losses.apply(
        lambda r: int(r["minutes"] * (2400 if r["line"] == "Line 1" else 2200) * 0.93), axis=1)

    spc_rows = []
    for r in spc_results:
        last = r["table"].tail(3)
        spc_rows.append({
            "line": r["line"], "characteristic": "Dome reversal pressure (psi)",
            "last_xbar": round(last["xbar"].iloc[-1], 2), "LCL": round(r["lcl"], 2),
            "CL": round(r["cl"], 2), "UCL": round(r["ucl"], 2),
            "Cpk_baseline": round(r["cpk_base"], 2), "Cpk_last_3wk": round(r["cpk_recent"], 2),
            "signal_last_3_shifts": "YES" if last["alert"].any() else "no",
        })
    spc_df = pd.DataFrame(spc_rows)

    actions = pd.DataFrame([
        {"#": 1, "issue": "Line 2 dome pressure trending down (SPC run rule)", "PDCA step": "Plan",
         "owner": "", "due": "", "status": "Open",
         "next step": "Check bodymaker dome tooling wear; compare with last tooling change date"},
        {"#": 2, "issue": "Line 2 decorator stops concentrated on Crew C", "PDCA step": "Plan",
         "owner": "", "due": "", "status": "Open",
         "next step": "Gemba with Crew C; review blanket/ink change SOP and training records"},
        {"#": 3, "issue": "Line 1 changeover time up since mid-August", "PDCA step": "Plan",
         "owner": "", "due": "", "status": "Open",
         "next step": "SMED study on next 3 changeovers; separate internal vs external tasks"},
    ])

    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        kpi.to_excel(xw, "Daily KPI", index=False)
        losses.to_excel(xw, "Top Losses", index=False)
        spc_df.to_excel(xw, "SPC Status", index=False)
        actions.to_excel(xw, "Action Log (PDCA)", index=False)

        ws = xw.sheets["Daily KPI"]
        style_sheet(ws, pct_cols=(4, 5, 6, 7, 8))
        ws.insert_rows(1)
        ws["A1"] = f"Daily Operations Meeting - {day:%d %b %Y}"
        ws["A1"].font = Font(bold=True, size=14)
        rng = f"G3:G{ws.max_row}"
        ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=[str(OEE_TARGET)], fill=RED_FILL))
        ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=[str(OEE_TARGET)], fill=GREEN_FILL))
        for row in ws.iter_rows(min_row=3, min_col=9, max_col=9):
            row[0].number_format = "+0.0;-0.0"
        for row in ws.iter_rows(min_row=3, min_col=10, max_col=10):
            row[0].number_format = "#,##0"
        ws.freeze_panes = "A3"

        style_sheet(xw.sheets["Top Losses"])
        for row in xw.sheets["Top Losses"].iter_rows(min_row=2, min_col=5, max_col=5):
            row[0].number_format = "#,##0"
        ws = xw.sheets["SPC Status"]
        style_sheet(ws)
        ws.conditional_formatting.add(f"I2:I{ws.max_row}",
                                      CellIsRule(operator="equal", formula=['"YES"'], fill=RED_FILL))
        style_sheet(xw.sheets["Action Log (PDCA)"], widths={2: 45, 7: 45})
    return day, kpi, losses


def write_insights(shifts, wk, paretos, crew, co, spc, dec_test, day, kpi):
    overall = oee_by(shifts, ["line"]).set_index("line")
    lines = ["# Key findings", "",
             f"Period: {shifts['date'].min():%d %b %Y} - {shifts['date'].max():%d %b %Y} "
             f"({shifts['date'].nunique()} days, 2 lines, 3 shifts/day)", "",
             "## Overall OEE", "",
             "| Line | Availability | Performance | Quality | OEE | vs target |",
             "|---|---|---|---|---|---|"]
    for line, r in overall.iterrows():
        lines.append(f"| {line} | {r.availability:.1%} | {r.performance:.1%} | {r.quality:.1%} | "
                     f"**{r.oee:.1%}** | {(r.oee - OEE_TARGET) * 100:+.1f} pts |")

    lines += ["", "## 1. Biggest losses (Pareto)", ""]
    for line, p in paretos.items():
        top = p[p["cum_pct"] <= 80].index.tolist() or [p.index[0]]
        lines.append(f"- **{line}:** {len(top)} of {len(p)} reasons cause ~80% of downtime: "
                     f"{', '.join(top)}.")

    lines += ["", "## 2. Line 2 decorator stops are crew-specific", "",
              f"- Average decorator downtime per shift: **Crew C {dec_test['crew_c_mean']:.1f} min** "
              f"vs other crews {dec_test['others_mean']:.1f} min "
              f"(Mann-Whitney U, p = {dec_test['p_value']:.1e}).",
              f"- Line 2 OEE by crew: " +
              ", ".join(f"{c} {v:.1%}" for c, v in crew.loc['Line 2'].items()) + ".",
              "- Same machine, different crews -> points to method/training (5M: Man/Method) "
              "rather than equipment. Recommended: gemba + SOP review with Crew C."]

    co["period"] = np.where(co["date"] < shifts["date"].min() + pd.Timedelta(days=45),
                            "first 45 days", "last 45 days")
    cmp = co.groupby("period")["minutes"].agg(["mean", "count", "sum"])
    t, p = stats.mannwhitneyu(co[co["period"] == "last 45 days"]["minutes"],
                              co[co["period"] == "first 45 days"]["minutes"], alternative="greater")
    lines += ["", "## 3. Line 1 changeovers got longer and more frequent", "",
              f"- First 45 days: {cmp.loc['first 45 days', 'count']:.0f} changeovers, "
              f"avg {cmp.loc['first 45 days', 'mean']:.0f} min.",
              f"- Last 45 days: {cmp.loc['last 45 days', 'count']:.0f} changeovers, "
              f"avg {cmp.loc['last 45 days', 'mean']:.0f} min (p = {p:.3f}).",
              f"- Extra changeover time: "
              f"{(cmp.loc['last 45 days', 'sum'] - cmp.loc['first 45 days', 'sum']) / 60:.0f} hours. "
              "Recommended: SMED study."]

    l2 = next(r for r in spc if r["line"] == "Line 2")
    lines += ["", "## 4. SPC early warning on Line 2 dome reversal pressure", "",
              f"- Control limits (first {SPC_BASELINE_DAYS} days): "
              f"LCL {l2['lcl']:.2f} / CL {l2['cl']:.2f} / UCL {l2['ucl']:.2f} psi.",
              f"- First low-side out-of-control signal after baseline: "
              f"**{l2['first_alert']:%d %b %Y}**. The first sampled can below LSL appeared on "
              f"{l2['first_out_of_spec']:%d %b %Y}, so the chart gave "
              f"**{(l2['first_out_of_spec'] - l2['first_alert']).days} days of early warning**.",
              f"- Cpk fell from **{l2['cpk_base']:.2f}** to **{l2['cpk_recent']:.2f}** in the last 3 weeks.",
              "- Gradual downward drift is typical of tooling wear -> act before LSL is reached."]

    dk = kpi[kpi["shift"] == "DAY TOTAL"]
    lines += ["", f"## Daily meeting snapshot - {day:%d %b %Y}", ""]
    for _, r in dk.iterrows():
        status = "on target" if r.oee >= OEE_TARGET else "below target"
        lines.append(f"- {r.line}: OEE {r.oee:.1%} ({status}, {r.gap_pts:+.1f} pts)")

    (OUT / "insights.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    return "\n".join(lines)


def main():
    CHARTS.mkdir(parents=True, exist_ok=True)
    shifts, events, quality = load()
    shifts = add_oee(shifts)

    wk = chart_weekly_oee(shifts)
    paretos = chart_pareto(events)
    crew = chart_crew_heatmap(shifts)
    co = chart_changeover(events)
    spc = [spc_analysis(quality, l) for l in ["Line 1", "Line 2"]]
    dec_test = decorator_crew_test(shifts, events)
    day, kpi, _ = daily_report(shifts, events, spc, OUT / "daily_ops_report.xlsx")
    print(write_insights(shifts, wk, paretos, crew, co, spc, dec_test, day, kpi))


if __name__ == "__main__":
    main()
